"""Small office files and images built in memory, for tests about formats and conversion."""

import io

from PIL import Image


def docx_bytes(*paragraphs: str) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs or ("Standard operating procedure for goods receipt.",):
        document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def pptx_bytes(title: str = "Quarterly supplier review") -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = title
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Batch", "Quantity"])
    sheet.append(["B-2041", 120])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def image_bytes(fmt: str, size: tuple[int, int] = (850, 1100), frames: int = 1) -> bytes:
    """A white page with a black bar; `frames` pages for a TIFF."""
    pages = []
    for n in range(frames):
        image = Image.new("RGB", size, "white")
        image.paste((0, 0, 0), (100, 100 + 40 * n, 700, 130 + 40 * n))
        pages.append(image)
    out = io.BytesIO()
    if frames > 1:
        pages[0].save(out, format=fmt, save_all=True, append_images=pages[1:])
    else:
        pages[0].save(out, format=fmt)
    return out.getvalue()


def zip_bytes() -> bytes:
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("notes.txt", "not an office file")
    return out.getvalue()
